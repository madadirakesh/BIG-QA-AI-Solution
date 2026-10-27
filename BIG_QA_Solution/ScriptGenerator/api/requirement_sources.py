"""
Requirement ingestion for the Test Case Generator.

Turns Jira tickets (description, acceptance criteria, comments, attachments) and
uploaded documents into clearly delimited, source-tagged text blocks. The blocks
go into the "What are we testing?" box, so the user sees exactly what the AI will
read, and the generation prompts can trace every statement back to its origin:

    ===== JIRA TICKET: ABC-12 | Story | Login page =====
    --- REQUIREMENT / DESCRIPTION ---
    ...
    ===== END JIRA TICKET: ABC-12 =====

Everything here is client-injected so it can be tested without Jira or an AI
provider: Jira calls go through the client passed in, image transcription goes
through the `image_reader` callable.
"""
from __future__ import annotations

import hashlib
import io
import os
import re
import threading
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from typing import Callable, Optional

# ── limits ───────────────────────────────────────────────────────────────────
MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_FILES_PER_UPLOAD = 10
MAX_DOCUMENT_CHARS = 200_000
MAX_SHEET_ROWS = 5_000
MAX_PDF_IMAGE_PAGES = 10          # scanned PDF pages sent to the AI for transcription
MAX_TICKETS_PER_IMPORT = 25
MAX_ATTACHMENTS_PER_TICKET = 10
MAX_COMMENTS_PER_TICKET = 100
MAX_DIRECT_CONTEXT_CHARS = 120_000  # larger contexts are condensed before generation
CONDENSE_CHUNK_CHARS = 40_000

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json"}
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | set(IMAGE_TYPES) | {".pdf", ".docx", ".xlsx", ".xls"}

ISSUE_KEY = re.compile(r"^[A-Z][A-Z0-9_]*-\d+$")
NO_CONTENT_MARKER = "NO_READABLE_CONTENT"

ImageReader = Callable[[bytes, str], str]

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"   # legacy Office and encrypted OOXML containers


class DocumentError(Exception):
    """A file that cannot be turned into text; the message is shown to the user as-is."""


@dataclass
class ExtractedDocument:
    name: str
    extension: str
    size: int
    text: str
    warnings: list = field(default_factory=list)


def human_size(n: int) -> str:
    if n < 1024:
        return f"{n} B"
    if n < 1024 * 1024:
        return f"{n / 1024:.0f} KB"
    return f"{n / (1024 * 1024):.1f} MB"


def _short(exc: Exception, limit: int = 160) -> str:
    text = str(exc).strip() or type(exc).__name__
    return text if len(text) <= limit else text[:limit] + "…"


# ── document extraction ──────────────────────────────────────────────────────

def extract_document(name: str, data: bytes, image_reader: Optional[ImageReader] = None) -> ExtractedDocument:
    """Validates a file and extracts its readable text, keeping page / sheet / heading references."""
    ext = os.path.splitext(name or "")[1].lower()
    if not data:
        raise DocumentError("The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError(f"The file is {human_size(len(data))}; the limit is {human_size(MAX_FILE_BYTES)}.")
    if ext == ".doc":
        raise DocumentError("Legacy .doc files are not supported. Save the file as .docx and upload it again.")
    if ext not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(e.lstrip(".").upper() for e in SUPPORTED_EXTENSIONS))
        raise DocumentError(f"Unsupported file type '{ext or 'none'}'. Supported types: {supported}.")

    warnings: list = []
    if ext == ".pdf":
        text = _pdf_text(data, warnings, image_reader)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext == ".xlsx":
        text = _xlsx_text(data, warnings)
    elif ext == ".xls":
        text = _xls_text(data, warnings)
    elif ext in IMAGE_TYPES:
        text = _image_text(data, ext, image_reader)
    else:
        text = _plain_text(data)

    text = _tidy(text)
    if not text:
        raise DocumentError("No readable text was found in the file.")
    if len(text) > MAX_DOCUMENT_CHARS:
        warnings.append(f"Only the first {MAX_DOCUMENT_CHARS:,} of {len(text):,} characters were extracted. "
                        "Split the document to include the rest.")
        text = text[:MAX_DOCUMENT_CHARS] + "\n[... extraction limit reached: the rest of this document is not included ...]"
    return ExtractedDocument(name=name, extension=ext.lstrip("."), size=len(data), text=text, warnings=warnings)


def _tidy(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _plain_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError as exc:
            raise DocumentError("The text file could not be decoded.") from exc
    if b"\x00" in data[:4096]:
        raise DocumentError("The file looks binary, not text. It may be corrupted or have the wrong extension.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _is_encryption_error(exc: BaseException) -> bool:
    seen = set()
    stack = [exc]
    while stack:
        e = stack.pop()
        if e is None or id(e) in seen:
            continue
        seen.add(id(e))
        if re.search(r"password|encrypt", f"{type(e).__name__} {e}", re.I):
            return True
        stack.extend([e.__cause__, e.__context__])
        stack.extend(a for a in getattr(e, "args", ()) if isinstance(a, BaseException))
    return False


def _pdf_text(data: bytes, warnings: list, image_reader: Optional[ImageReader]) -> str:
    if not data.lstrip()[:5].startswith(b"%PDF"):
        raise DocumentError("The file is not a valid PDF (it may be corrupted or renamed).")
    try:
        import pdfplumber
    except ImportError as exc:
        raise DocumentError("PDF support needs the 'pdfplumber' package on the server (pip install pdfplumber).") from exc

    parts, blank, transcribed = [], [], 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for number, page in enumerate(pdf.pages, 1):
                text = (page.extract_text() or "").strip()
                if not text and image_reader and transcribed < MAX_PDF_IMAGE_PAGES:
                    png = _render_pdf_page(page)
                    if png:
                        transcribed += 1
                        try:
                            text = image_reader(png, "image/png").strip()
                        except Exception as exc:
                            warnings.append(f"Page {number}: the AI provider could not read the scanned page ({_short(exc)}).")
                            text = ""
                        if text.upper() == NO_CONTENT_MARKER:
                            text = ""
                if text:
                    parts.append(f"[Page {number}]\n{text}")
                else:
                    blank.append(number)
    except DocumentError:
        raise
    except Exception as exc:
        if _is_encryption_error(exc):
            raise DocumentError("The PDF is password-protected. Remove the password and upload it again.") from exc
        raise DocumentError(f"The PDF could not be read; it may be corrupted ({_short(exc)}).") from exc

    if blank and parts:
        warnings.append(f"No text found on page(s) {', '.join(map(str, blank[:20]))}"
                        f"{'…' if len(blank) > 20 else ''} (blank or scanned pages).")
    if not parts and blank and not image_reader:
        raise DocumentError("The PDF has no extractable text (it looks scanned) and image reading is not available.")
    return "\n\n".join(parts)


def _render_pdf_page(page) -> Optional[bytes]:
    try:
        image = page.to_image(resolution=150).original
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        return buf.getvalue()
    except Exception:
        return None


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_text(data: bytes) -> str:
    if data.startswith(_OLE_MAGIC):
        raise DocumentError("The Word document is password-protected (encrypted) or in legacy .doc format. "
                            "Remove the password / save as .docx and upload it again.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            xml = zf.read("word/document.xml")
    except KeyError as exc:
        raise DocumentError("The file is not a valid Word document (word/document.xml is missing).") from exc
    except zipfile.BadZipFile as exc:
        raise DocumentError("The Word document is corrupted or is not a .docx file.") from exc
    try:
        root = ET.fromstring(xml)
    except ET.ParseError as exc:
        raise DocumentError("The Word document is corrupted (its content could not be parsed).") from exc

    body = root.find(f"{_W}body")
    lines = []
    for el in (body if body is not None else []):
        if el.tag == f"{_W}p":
            line = _docx_paragraph(el)
            if line:
                lines.append(line)
        elif el.tag == f"{_W}tbl":
            rows = _docx_table(el)
            if rows:
                lines.append("[Table]")
                lines.extend(rows)
                lines.append("")
    return "\n".join(lines)


def _docx_run_text(el) -> str:
    out = []
    for node in el.iter():
        if node.tag == f"{_W}t":
            out.append(node.text or "")
        elif node.tag == f"{_W}tab":
            out.append("\t")
        elif node.tag in (f"{_W}br", f"{_W}cr"):
            out.append("\n")
    return "".join(out)


def _docx_paragraph(p) -> str:
    text = _docx_run_text(p).strip()
    if not text:
        return ""
    ppr = p.find(f"{_W}pPr")
    style = ""
    if ppr is not None:
        style_el = ppr.find(f"{_W}pStyle")
        if style_el is not None:
            style = style_el.get(f"{_W}val", "")
        if ppr.find(f"{_W}numPr") is not None:
            return f"- {text}"
    heading = re.match(r"heading\s*(\d)", style, re.I)
    if heading:
        return f"{'#' * min(int(heading.group(1)), 6)} {text}"
    if style.lower() == "title":
        return f"# {text}"
    return text


def _docx_table(tbl) -> list:
    rows = []
    for tr in tbl.findall(f"{_W}tr"):
        cells = []
        for tc in tr.findall(f"{_W}tc"):
            paragraphs = [_docx_run_text(p).strip() for p in tc.iter(f"{_W}p")]
            cells.append(" ".join(t for t in paragraphs if t))
        if any(cells):
            rows.append(" | ".join(cells))
    return rows


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\n", " ").strip()


def _sheet_rows(title: str, rows_iter, warnings: list) -> str:
    rows = []
    for number, row in enumerate(rows_iter, 1):
        cells = [_cell_text(v) for v in row]
        while cells and not cells[-1]:
            cells.pop()
        if not any(cells):
            continue
        if len(rows) >= MAX_SHEET_ROWS:
            warnings.append(f"Sheet '{title}': only the first {MAX_SHEET_ROWS:,} non-empty rows were extracted.")
            break
        rows.append(f"Row {number}: " + " | ".join(cells))
    return f"[Sheet: {title}]\n" + "\n".join(rows) if rows else ""


def _xlsx_text(data: bytes, warnings: list) -> str:
    if data.startswith(_OLE_MAGIC):
        raise DocumentError("The workbook is password-protected (encrypted) or in legacy .xls format. "
                            "Remove the password / save as .xlsx and upload it again.")
    try:
        import openpyxl
    except ImportError as exc:
        raise DocumentError("Excel support needs the 'openpyxl' package on the server.") from exc
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentError(f"The workbook is corrupted or not a valid .xlsx file ({_short(exc)}).") from exc
    try:
        parts = [_sheet_rows(ws.title, ws.iter_rows(values_only=True), warnings) for ws in wb.worksheets]
    finally:
        wb.close()
    return "\n\n".join(p for p in parts if p)


def _xls_text(data: bytes, warnings: list) -> str:
    try:
        import xlrd
    except ImportError as exc:
        raise DocumentError("Reading .xls files needs the 'xlrd' package, which is not installed on the server. "
                            "Save the workbook as .xlsx and upload it again.") from exc
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        if _is_encryption_error(exc):
            raise DocumentError("The workbook is password-protected. Remove the password and upload it again.") from exc
        raise DocumentError(f"The workbook is corrupted or not a valid .xls file ({_short(exc)}).") from exc
    parts = []
    for sheet in book.sheets():
        rows = (sheet.row_values(r) for r in range(sheet.nrows))
        parts.append(_sheet_rows(sheet.name, rows, warnings))
    return "\n\n".join(p for p in parts if p)


def _image_text(data: bytes, ext: str, image_reader: Optional[ImageReader]) -> str:
    valid = data.startswith(b"\x89PNG\r\n\x1a\n") if ext == ".png" else data.startswith(b"\xff\xd8\xff")
    if not valid:
        raise DocumentError("The image is corrupted or its content does not match its extension.")
    try:
        from PIL import Image
        Image.open(io.BytesIO(data)).verify()
    except ImportError:
        pass
    except Exception as exc:
        raise DocumentError(f"The image is corrupted ({_short(exc)}).") from exc
    if image_reader is None:
        raise DocumentError("Reading images needs the configured AI provider, which is not available.")
    try:
        text = (image_reader(data, IMAGE_TYPES[ext]) or "").strip()
    except Exception as exc:
        raise DocumentError(f"The AI provider could not read the image ({_short(exc)}). "
                            "Check that the configured model supports images.") from exc
    if not text or text.upper() == NO_CONTENT_MARKER:
        raise DocumentError("No readable text or requirement content was found in the image.")
    return text


# ── Jira field normalisation ─────────────────────────────────────────────────

def _holder_to_dict(value):
    """jira's PropertyHolder (used for JSON objects it does not model) -> plain dict."""
    if isinstance(value, dict):
        return {k: _holder_to_dict(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_holder_to_dict(v) for v in value]
    if hasattr(value, "__dict__") and not isinstance(value, (str, bytes)):
        return {k: _holder_to_dict(v) for k, v in vars(value).items() if not k.startswith("_")}
    return value


def to_plain_text(value) -> str:
    """Normalises a Jira field value (wiki markup string, ADF document, option, list...) to readable text."""
    if value is None:
        return ""
    if isinstance(value, str):
        return clean_jira_markup(value)
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, (list, tuple)):
        return "\n".join(t for t in (to_plain_text(v) for v in value) if t)
    if isinstance(value, dict):
        if value.get("type") == "doc" or isinstance(value.get("content"), list):
            return _tidy(_adf_to_text(value))
        for key in ("value", "name", "displayName"):
            if key in value:
                return to_plain_text(value[key])
        return ""
    for attr in ("value", "name", "displayName"):
        if isinstance(getattr(value, attr, None), str):
            return clean_jira_markup(getattr(value, attr))
    if hasattr(value, "__dict__"):
        return to_plain_text(_holder_to_dict(value))
    return str(value).strip()


_ADF_BLOCKS = {"paragraph", "heading", "blockquote", "codeBlock", "panel", "rule", "mediaSingle", "mediaGroup"}


def _adf_to_text(node) -> str:
    if not isinstance(node, dict):
        return ""
    kind = node.get("type")
    attrs = node.get("attrs") or {}
    children = node.get("content") or []
    if kind == "text":
        return node.get("text", "")
    if kind == "hardBreak":
        return "\n"
    if kind == "mention":
        return attrs.get("text", "")
    if kind == "emoji":
        return attrs.get("text") or attrs.get("shortName", "")
    if kind in ("inlineCard", "blockCard"):
        return attrs.get("url", "")
    if kind == "media":
        return f"[attachment: {attrs.get('alt') or attrs.get('id', '')}]"
    if kind == "bulletList":
        return "\n".join("- " + _adf_to_text(c).strip() for c in children) + "\n"
    if kind == "orderedList":
        start = int(attrs.get("order", 1) or 1)
        return "\n".join(f"{start + i}. " + _adf_to_text(c).strip() for i, c in enumerate(children)) + "\n"
    if kind == "listItem":
        return "\n  ".join(_adf_to_text(c).strip() for c in children)
    if kind == "table":
        return "\n".join(_adf_to_text(row) for row in children) + "\n\n"
    if kind == "tableRow":
        return " | ".join(_adf_to_text(cell).strip() for cell in children)
    inner = "".join(_adf_to_text(c) for c in children)
    if kind == "heading":
        return "#" * int(attrs.get("level", 1) or 1) + " " + inner.strip() + "\n"
    if kind in _ADF_BLOCKS:
        return inner + "\n"
    return inner


def clean_jira_markup(text: str) -> str:
    """Light clean-up of Jira wiki markup so it reads as plain text without losing content."""
    t = (text or "").replace("\r\n", "\n")
    t = re.sub(r"(?m)^\s*h([1-6])\.\s*", lambda m: "#" * int(m.group(1)) + " ", t)
    t = re.sub(r"\{color(?::[^}]*)?\}", "", t)
    t = re.sub(r"\{(?:noformat|code|panel|quote)(?::[^}]*)?\}", "", t)
    t = re.sub(r"!([^!|\n]+)(?:\|[^!\n]*)?!", r"[image: \1]", t)
    t = re.sub(r"\[([^|\]\n]+)\|([^\]\n]+)\]", r"\1 (\2)", t)
    return _tidy(t)


_AC_HEADING = re.compile(
    r"^\s*(?:h[1-6]\.\s*|#{1,6}\s*)?[*_]*\s*acceptance\s+criteria\s*[*_]*\s*(?P<colon>[:\-–])?\s*[*_]*\s*(?P<rest>.*?)\s*$",
    re.I)
_SECTION_END = re.compile(
    r"^\s*(?:h[1-6]\.\s+|#{1,6}\s+\S|\*[^*\n]{2,60}:?\*\s*:?\s*$|"
    r"(?:notes?|out of scope|dependencies|assumptions|technical notes|definition of done|description|"
    r"background|design|attachments?|references?|open questions)\s*:\s*$)", re.I)


def split_acceptance_criteria(description: str) -> tuple:
    """Finds an 'Acceptance Criteria' section in a description. Returns (criteria, description_without_it)."""
    lines = (description or "").split("\n")
    for i, line in enumerate(lines):
        m = _AC_HEADING.match(line)
        if not m or (m.group("rest") and not m.group("colon")):
            continue
        collected = [m.group("rest").strip(" *_")] if m.group("rest") else []
        end = len(lines)
        for j in range(i + 1, len(lines)):
            if _SECTION_END.match(lines[j]) and not _AC_HEADING.match(lines[j]):
                end = j
                break
            collected.append(lines[j])
        criteria = _tidy("\n".join(collected))
        if criteria:
            remaining = _tidy("\n".join(lines[:i] + lines[end:]))
            return criteria, remaining
    return "", description or ""


def acceptance_criteria_field_ids(jira) -> list:
    """(field id, field name) of every Jira field named like 'Acceptance Criteria'."""
    try:
        fields = jira.fields()
    except Exception:
        return []
    return [(f["id"], f.get("name", f["id"])) for f in fields
            if "acceptance criteria" in (f.get("name") or "").lower()]


def describe_jira_error(exc: Exception, key: str) -> str:
    status = getattr(exc, "status_code", None)
    if status == 404:
        return f"{key} was not found, or you do not have permission to view it."
    if status == 401:
        return "Jira rejected the credentials. Check the Jira settings under Configurations."
    if status == 403:
        return f"You do not have permission to view {key}."
    text = getattr(exc, "text", "") or str(exc)
    return f"Could not load {key}: {_short(Exception(text))}"


class AttachmentRegistry:
    """Remembers attachment contents already imported in this run, so the same file is included once."""

    def __init__(self):
        self._seen = {}
        self._lock = threading.Lock()

    def claim(self, digest: str, label: str) -> Optional[str]:
        """Returns the label of the earlier copy, or None if this is the first time the content is seen."""
        with self._lock:
            if digest in self._seen:
                return self._seen[digest]
            self._seen[digest] = label
            return None


def _import_attachment(att, key: str, image_reader, registry: AttachmentRegistry) -> dict:
    name = str(getattr(att, "filename", "") or "attachment")
    size = int(getattr(att, "size", 0) or 0)
    entry = {"name": name, "size": size, "status": "", "message": "", "text": "", "warnings": []}
    ext = os.path.splitext(name)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        entry.update(status="skipped", message=f"Unsupported file type ({ext or 'no extension'}).")
        return entry
    if size > MAX_FILE_BYTES:
        entry.update(status="skipped", message=f"Too large ({human_size(size)}; limit {human_size(MAX_FILE_BYTES)}).")
        return entry
    try:
        data = att.get()
    except Exception as exc:
        entry.update(status="failed", message=f"Could not download it (check attachment permissions): {_short(exc)}")
        return entry
    earlier = registry.claim(hashlib.sha256(data or b"").hexdigest(), f"{key} / {name}")
    if earlier:
        entry.update(status="duplicate", message=f"Same content as {earlier}; included only once.")
        return entry
    try:
        doc = extract_document(name, data, image_reader)
    except DocumentError as exc:
        entry.update(status="failed", message=str(exc))
        return entry
    entry.update(status="extracted", text=doc.text, warnings=doc.warnings,
                 message=f"{len(doc.text):,} characters extracted.")
    return entry


def import_ticket(jira, key: str, include: dict, ac_fields: list, image_reader: Optional[ImageReader] = None,
                  registry: Optional[AttachmentRegistry] = None) -> dict:
    """
    Retrieves only the requested information types for one ticket. A missing field
    becomes a warning; only a ticket that cannot be loaded at all is a failure.
    """
    registry = registry or AttachmentRegistry()
    result = {"key": key, "ok": False, "summary": "", "type": "", "status": "", "warnings": [], "error": ""}
    want_desc = bool(include.get("description"))
    want_ac = bool(include.get("acceptance_criteria"))
    want_comments = bool(include.get("comments"))
    want_atts = bool(include.get("attachments"))

    fields = ["summary", "issuetype", "status"]
    if want_desc or want_ac:
        fields.append("description")   # acceptance criteria may live in the description
    if want_ac:
        fields.extend(fid for fid, _ in ac_fields)
    if want_comments:
        fields.append("comment")
    if want_atts:
        fields.append("attachment")
    try:
        issue = jira.issue(key, fields=",".join(fields))
    except Exception as exc:
        result["error"] = describe_jira_error(exc, key)
        return result

    f = issue.fields
    result.update(key=issue.key, summary=str(getattr(f, "summary", "") or ""),
                  type=str(getattr(f, "issuetype", "") or ""), status=str(getattr(f, "status", "") or ""))
    description = to_plain_text(getattr(f, "description", None))
    warnings = result["warnings"]

    if want_ac:
        criteria, source = "", ""
        for fid, fname in ac_fields:
            criteria = to_plain_text(getattr(f, fid, None))
            if criteria:
                source = f"Jira field '{fname}'"
                break
        if not criteria:
            criteria, without = split_acceptance_criteria(description)
            if criteria:
                source = "'Acceptance Criteria' section of the description"
                description = without   # shown once, under Acceptance Criteria
        if not criteria:
            warnings.append("No acceptance criteria found (no Acceptance Criteria field value and no such "
                            "section in the description).")
        result["acceptance_criteria"] = criteria
        result["acceptance_criteria_source"] = source

    if want_desc:
        result["description"] = description
        if not description:
            warnings.append("The ticket has no description.")

    if want_comments:
        holder = getattr(f, "comment", None)
        comments = list(getattr(holder, "comments", None) or [])
        total = int(getattr(holder, "total", len(comments)) or len(comments))
        if total > len(comments):
            warnings.append(f"Jira returned {len(comments)} of {total} comments.")
        if len(comments) > MAX_COMMENTS_PER_TICKET:
            warnings.append(f"Only the latest {MAX_COMMENTS_PER_TICKET} of {len(comments)} comments were included.")
            comments = comments[-MAX_COMMENTS_PER_TICKET:]
        rows = []
        for c in comments:
            body = to_plain_text(getattr(c, "body", ""))
            if body:
                author = getattr(c, "author", None)
                rows.append({"author": str(getattr(author, "displayName", "") or "Unknown"),
                             "created": str(getattr(c, "created", "") or "")[:16].replace("T", " "),
                             "body": body})
        if not rows:
            warnings.append("The ticket has no comments.")
        result["comments"] = rows

    if want_atts:
        attachments = list(getattr(f, "attachment", None) or [])
        if len(attachments) > MAX_ATTACHMENTS_PER_TICKET:
            warnings.append(f"Only the first {MAX_ATTACHMENTS_PER_TICKET} of {len(attachments)} attachments were processed.")
        result["attachments"] = [_import_attachment(a, result["key"], image_reader, registry)
                                 for a in attachments[:MAX_ATTACHMENTS_PER_TICKET]]
        if not attachments:
            warnings.append("The ticket has no attachments.")
        for a in result["attachments"]:
            if a["status"] in ("failed", "skipped"):
                warnings.append(f"Attachment '{a['name']}' not included: {a['message']}")
            warnings.extend(f"Attachment '{a['name']}': {w}" for w in a["warnings"])

    result["ok"] = True
    result["block"] = render_ticket_block(result, include)
    return result


# ── source blocks ────────────────────────────────────────────────────────────

def neutralize_markers_in_prompt(text: str) -> str:
    """The requirement text must not be able to close (or reopen) the prompt's <requirements_data> section."""
    return re.sub(r"</?\s*requirements_data\s*>", "[requirements_data tag removed]", text or "", flags=re.I)


def neutralize_markers(text: str) -> str:
    """Imported content must not be able to fake a source boundary or close the prompt's data section."""
    return neutralize_markers_in_prompt(re.sub(r"(?m)^(\s*)={5,}", r"\1= = = = =", text or ""))


def _one_line(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("=====", "=")).strip()


def render_ticket_block(t: dict, include: dict) -> str:
    key = t["key"]
    out = [f"===== JIRA TICKET: {key} | {_one_line(t.get('type')) or 'Issue'} | {_one_line(t.get('summary'))} =====",
           f"Status: {_one_line(t.get('status')) or 'Unknown'}"]

    def section(title: str, body: str):
        out.extend(["", f"--- {title} ---", neutralize_markers(body)])

    if include.get("description"):
        section("REQUIREMENT / DESCRIPTION", t.get("description") or "(No description in this ticket.)")
    if include.get("acceptance_criteria"):
        source = t.get("acceptance_criteria_source")
        section(f"ACCEPTANCE CRITERIA (source: {source})" if source else "ACCEPTANCE CRITERIA",
                t.get("acceptance_criteria") or "(Not available: this ticket has no acceptance criteria.)")
    if include.get("comments"):
        comments = t.get("comments") or []
        section(f"COMMENTS ({len(comments)})",
                "\n\n".join(f"[{c['created']}] {_one_line(c['author'])}:\n{c['body']}" for c in comments)
                or "(No comments.)")
    if include.get("attachments"):
        attachments = t.get("attachments") or []
        for a in attachments:
            if a["status"] == "extracted":
                section(f"ATTACHMENT: {_one_line(a['name'])} ({human_size(a['size'])})", a["text"])
        missing = [a for a in attachments if a["status"] != "extracted"]
        if missing:
            section("ATTACHMENTS NOT INCLUDED",
                    "\n".join(f"- {_one_line(a['name'])}: {a['message']}" for a in missing))
        if not attachments:
            section("ATTACHMENTS", "(No attachments.)")
    out.extend(["", f"===== END JIRA TICKET: {key} ====="])
    return "\n".join(out)


def render_document_block(doc: ExtractedDocument) -> str:
    name = _one_line(doc.name).replace("|", "/")
    return "\n".join([
        f"===== SUPPORTING DOCUMENT: {name} | {doc.extension.upper()} | {human_size(doc.size)} =====",
        neutralize_markers(doc.text),
        f"===== END SUPPORTING DOCUMENT: {name} =====",
    ])


# ── context size management ──────────────────────────────────────────────────

_BLOCK_START = re.compile(r"(?m)^===== (?:JIRA TICKET|SUPPORTING DOCUMENT): ")


def split_into_chunks(text: str, max_chars: int = CONDENSE_CHUNK_CHARS) -> list:
    """Splits the context at source-block boundaries; a block bigger than a chunk is split by lines."""
    starts = sorted({0, len(text), *(m.start() for m in _BLOCK_START.finditer(text))})
    pieces = []
    for a, b in zip(starts, starts[1:]):
        segment = text[a:b]
        if not segment.strip():
            continue
        pieces.extend([segment] if len(segment) <= max_chars else _split_segment(segment, max_chars))

    chunks, current = [], ""
    for piece in pieces:
        if current and len(current) + len(piece) > max_chars:
            chunks.append(current)
            current = ""
        current += piece
    if current.strip():
        chunks.append(current)
    return chunks


def _split_segment(segment: str, max_chars: int) -> list:
    header = segment.split("\n", 1)[0] if segment.startswith("=====") else ""
    prefix = f"[continued: {header}]\n" if header else ""
    budget = max(1000, max_chars - len(prefix))
    parts, current = [], ""
    for line in segment.splitlines(keepends=True):
        while len(line) > budget:   # a single huge line (e.g. minified text)
            if current:
                parts.append(current)
                current = ""
            parts.append(line[:budget])
            line = line[budget:]
        if len(current) + len(line) > budget:
            parts.append(current)
            current = ""
        current += line
    if current:
        parts.append(current)
    return [parts[0]] + [prefix + p for p in parts[1:]] if parts else []
