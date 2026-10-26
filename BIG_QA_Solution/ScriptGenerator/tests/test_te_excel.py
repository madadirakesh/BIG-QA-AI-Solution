import io

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from conftest import png_bytes
from test_execution import excel_io
from test_execution.models import TestCaseResult, StepResult, Evidence, PASSED, FAILED, compute_summary
from test_execution.security import PolicyError, validate_upload


def _xlsx(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SINGLE = [
    ["Test Case ID", "Test Case Title", "Preconditions", "Test Steps", "Test Data", "Expected Result", "Status"],
    ["TC-01", "Valid login", "User exists", "1. Open the application\n2. Enter valid username\n3. Click Login button",
     "", "1. Login page shown\n2. -\n3. 'Welcome back' is displayed", ""],
    ["TC-02", "Empty steps", "", "", "", "", ""],
    ["", "No id", "", "1. Open the application", "", "", ""],
    ["TC-01", "Duplicate", "", "1. Open the application", "", "", ""],
]

# Exactly what the Test Case Generator exports in "separate rows" mode.
GENERATOR = [
    ["ID", "Work Item Type", "Title", "State", "Test Step", "Step Action", "Step Expected"],
    ["TC-001", "Test Case", "Login works", "Design", "", "", ""],
    ["", "", "", "", 1, "Navigate to the login page", "Login page is displayed"],
    ["", "", "", "", 2, "Enter 'alice' in Username field", ""],
    ["", "", "", "", 3, "Click on Login button", "'Welcome back' message is displayed"],
    ["TC-002", "Test Case", "Logout", "Design", "", "", ""],
    ["", "", "", "", 1, "Click Logout link", "'Sign in' is displayed"],
]


def test_mapping_detects_generator_layout():
    sheet = excel_io.read_workbook(_xlsx({"Test Cases": GENERATOR}), ".xlsx")[0]
    m = excel_io.suggest_mapping(sheet.headers)
    assert m["id"] == 0 and m["title"] == 2
    assert m["step_number"] == 4 and m["steps"] == 5 and m["step_expected"] == 6


def test_separate_rows_layout_groups_steps_under_metadata_row():
    sheet = excel_io.read_workbook(_xlsx({"Test Cases": GENERATOR}), ".xlsx")[0]
    cases = excel_io.parse_test_cases(sheet, excel_io.suggest_mapping(sheet.headers))
    assert [c.key for c in cases] == ["TC-001", "TC-002"]
    tc = cases[0]
    assert [s.action_text for s in tc.steps] == ["Navigate to the login page", "Enter 'alice' in Username field", "Click on Login button"]
    assert tc.steps[2].expected == "'Welcome back' message is displayed"
    assert tc.source_ref["row"] == 2 and tc.source_ref["step_rows"] == [3, 4, 5]
    assert all(c.valid for c in cases)


def test_single_cell_layout_and_validation_never_drops_rows():
    sheet = excel_io.read_workbook(_xlsx({"Sheet1": SINGLE}), ".xlsx")[0]
    cases = excel_io.parse_test_cases(sheet, excel_io.suggest_mapping(sheet.headers))
    assert len(cases) == 4                      # every row is accounted for
    tc1 = cases[0]
    assert [s.number for s in tc1.steps] == [1, 2, 3]
    assert tc1.steps[2].expected == "'Welcome back' is displayed"   # numbered expectations aligned
    assert tc1.preconditions == "User exists"
    assert "No test steps" in cases[1].issues
    assert "Missing Test Case ID" in cases[2].issues and cases[2].key == "ROW-4"
    assert any("Duplicate" in i for i in cases[3].issues) and cases[3].key == "TC-01#2"


def test_header_row_offset_and_multiple_sheets():
    data = _xlsx({"Cover": [["Release 1.0 test plan"]], "Cases": [["Project X"], []] + SINGLE})
    sheets = excel_io.read_workbook(data, ".xlsx")
    assert [s.name for s in sheets] == ["Cover", "Cases"]
    cases_sheet = sheets[1]
    assert cases_sheet.header_row == 3
    assert excel_io.suggest_mapping(cases_sheet.headers)["steps"] == 3


def test_split_numbered_variants():
    assert excel_io.split_numbered("1. a\n2) b\n3 - c") == [(1, "a"), (2, "b"), (3, "c")]
    assert excel_io.split_numbered("Step 1: open\nStep 2: click") == [(1, "open"), (2, "click")]
    assert excel_io.split_numbered("- a\n- b") == [(None, "a"), (None, "b")]
    assert excel_io.split_numbered("just one step") == [(None, "just one step")]


def test_upload_validation():
    good = _xlsx({"S": [["a", "b"]]})
    assert validate_upload("x.xlsx", good) == ".xlsx"
    with pytest.raises(PolicyError):
        validate_upload("x.csv", b"a,b")
    with pytest.raises(PolicyError):
        validate_upload("x.xlsx", b"not a zip")
    with pytest.raises(PolicyError):
        validate_upload("x.xls", good)            # extension/content mismatch
    with pytest.raises(PolicyError):
        validate_upload("x.xlsx", b"")


def _results(cases, with_evidence=True):
    res = {}
    tc = cases[0]
    r = TestCaseResult(run_id="TE-1", test_key=tc.key, title=tc.title, source="excel", status=FAILED,
                       started_at="2026-10-23T10:00:00+00:00", ended_at="2026-10-23T10:00:05+00:00", duration_ms=5000,
                       failed_step=3, failure_reason='Expected text "Welcome back" was not found', actual="boom",
                       steps=[StepResult(1, "a", "", PASSED, "ok"), StepResult(2, "b", "", PASSED, "ok"),
                              StepResult(3, "c", "", FAILED, "not found")])
    if with_evidence:
        r.evidence = [Evidence("ev1", tc.key, 3, FAILED, "2026-10-23T10:00:05+00:00", "TE-1/a.png", 10),
                      Evidence("ev2", tc.key, None, FAILED, "2026-10-23T10:00:05+00:00", "TE-1/b.png", 10)]
    res[tc.key] = r
    res[cases[1].key] = TestCaseResult(run_id="TE-1", test_key=cases[1].key, title="", source="excel", status=PASSED)
    return res


def test_result_workbook_preserves_original_and_writes_results():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    for r in GENERATOR:
        ws.append(r)
    ws["C2"].fill = PatternFill("solid", fgColor="FFFF00")
    other = wb.create_sheet("Notes")
    other["A1"] = "keep me"
    buf = io.BytesIO()
    wb.save(buf)
    original = buf.getvalue()

    sheet = excel_io.read_workbook(original, ".xlsx")[0]
    mapping = excel_io.suggest_mapping(sheet.headers)
    cases = excel_io.parse_test_cases(sheet, mapping)
    results = _results(cases)
    summary = compute_summary(list(results.values()))
    out, notes = excel_io.build_result_workbook(
        original, ".xlsx", "Test Cases", sheet.header_row, mapping, cases, results,
        evidence_reader=lambda ev: png_bytes(ev.evidence_id), evidence_url=lambda ev: f"http://h/{ev.evidence_id}",
        embed=True, summary={**summary, "run_id": "TE-1"})
    wb2 = openpyxl.load_workbook(io.BytesIO(out))
    assert wb2["Notes"]["A1"].value == "keep me"
    ws2 = wb2["Test Cases"]
    assert ws2["C2"].fill.fgColor.rgb.endswith("FFFF00")          # original formatting kept
    headers = [c.value for c in ws2[1]]
    col = {h: i + 1 for i, h in enumerate(headers)}
    assert ws2.cell(2, col["Execution Status"]).value == FAILED
    assert "Step 3" in ws2.cell(2, col["Failure Details"]).value
    assert ws2.cell(5, col["Execution Status"]).value == FAILED     # per-step row for step 3
    assert ws2.cell(3, col["Execution Status"]).value == PASSED
    assert ws2.cell(6, col["Execution Status"]).value == PASSED     # TC-002
    ev_cell = ws2.cell(2, col["Evidence"])
    assert "2 screenshot" in ev_cell.value and ev_cell.hyperlink is not None
    ev_ws = wb2[excel_io.EVIDENCE_SHEET]
    assert len(ev_ws._images) == 2
    assert ev_ws["B2"].value == "Step 3" and ev_ws["B3"].value == "Final"
    assert wb2[excel_io.SUMMARY_SHEET]["B4"].value == 1             # Passed


def test_result_workbook_links_instead_of_embedding():
    original = _xlsx({"Test Cases": GENERATOR})
    sheet = excel_io.read_workbook(original, ".xlsx")[0]
    mapping = excel_io.suggest_mapping(sheet.headers)
    cases = excel_io.parse_test_cases(sheet, mapping)
    out, _ = excel_io.build_result_workbook(
        original, ".xlsx", "Test Cases", sheet.header_row, mapping, cases, _results(cases),
        evidence_reader=lambda ev: png_bytes("x"), evidence_url=lambda ev: f"http://h/{ev.evidence_id}",
        embed=False, summary={})
    ev_ws = openpyxl.load_workbook(io.BytesIO(out))[excel_io.EVIDENCE_SHEET]
    assert not ev_ws._images
    assert ev_ws["E2"].hyperlink.target == "http://h/ev1"


def test_existing_status_and_actual_columns_are_reused():
    original = _xlsx({"S": SINGLE})
    sheet = excel_io.read_workbook(original, ".xlsx")[0]
    mapping = excel_io.suggest_mapping(sheet.headers)
    cases = excel_io.parse_test_cases(sheet, mapping)
    out, _ = excel_io.build_result_workbook(original, ".xlsx", "S", 1, mapping, cases, _results(cases, False),
                                            lambda e: None, lambda e: "", True, {})
    ws = openpyxl.load_workbook(io.BytesIO(out))["S"]
    assert ws.cell(2, 7).value == FAILED        # the existing "Status" column (G)
    headers = [c.value for c in ws[1]]
    assert headers.count("Execution Status") == 0
