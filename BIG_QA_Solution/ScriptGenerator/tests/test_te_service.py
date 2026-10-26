"""End-to-end through TestExecutionService with a fake Playwright MCP and fake Jira (mocked, not live)."""
import io
import json
import time

import openpyxl
import pytest

from conftest import FakeMCPSession
from test_te_integrations import FakeJira, _issue
from test_execution.models import RunConfig, PUB_PUBLISHED, PUB_FAILED
from test_execution.security import PolicyError
from test_execution.service import TestExecutionService, RUN_COMPLETED, RUN_INTERRUPTED
from test_execution.store import ExecutionStore

BASE = "http://app.test/login"
CREDS = {"username": "alice", "password": "s3cret!"}

ROWS = [
    ["ID", "Title", "Test Step", "Step Action", "Step Expected"],
    ["TC-001", "Login works", "", "", ""],
    ["", "", 1, "Open the application", "Page title is 'Demo Login'"],
    ["", "", 2, "Enter valid username", ""],
    ["", "", 3, "Enter valid password", ""],
    ["", "", 4, "Click on Login button", "'Welcome back' is displayed"],
    ["TC-002", "Wrong password fails", "", "", ""],
    ["", "", 1, "Open the application", ""],
    ["", "", 2, "Enter 'alice' in Username field", ""],
    ["", "", 3, "Enter 'nope' in Password field", ""],
    ["", "", 4, "Click on Login button", "'Welcome back' is displayed"],
    ["TC-003", "No steps", "", "", ""],
]


def _xlsx(rows):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def jira():
    j = FakeJira()
    j.issues["PROJ-7"] = _issue("PROJ-7", "Test", "Login", "Steps:\n# Open the application\n# Enter valid username\n"
                                "# Enter valid password\n# Click on Login button -> 'Welcome back' is displayed")
    return j


@pytest.fixture
def svc(tmp_path, jira):
    return TestExecutionService(str(tmp_path / "data"), str(tmp_path / "db.sqlite"), jira_client_factory=lambda: jira,
                                ai_call=None, session_factory=lambda i: FakeMCPSession())


def wait(svc, run_id, timeout=30):
    end = time.time() + timeout
    while svc.is_active(run_id):
        assert time.time() < end, "run did not finish"
        time.sleep(0.05)
    return svc.run_view(run_id)


def cfg(**kw):
    return RunConfig.from_dict({"base_url": BASE, "screenshot_mode": "final", "action_timeout_ms": 1000, **kw})


def test_excel_end_to_end(svc):
    up = svc.save_upload("qa", "cases.xlsx", _xlsx(ROWS))
    sheet = up["sheets"][0]
    cases, meta = svc.excel_cases(up["upload_id"], sheet["name"], sheet["suggested_mapping"])
    prev = svc.preview(cases, True, True)
    assert prev["stats"]["total"] == 3 and prev["stats"]["invalid"] == 1 and prev["stats"]["steps_needing_ai"] == 0

    run_id = svc.start_run("qa", "excel", cases, cfg(), meta, CREDS)
    view = wait(svc, run_id)
    assert view["run"]["status"] == RUN_COMPLETED
    by = {r["test_key"]: r for r in view["results"]}
    assert by["TC-001"]["status"] == "Passed"
    assert by["TC-002"]["status"] == "Failed" and by["TC-002"]["failed_step"] == 4
    assert by["TC-003"]["status"] == "Blocked"
    s = view["summary"]
    assert (s["passed"], s["failed"], s["blocked"]) == (1, 1, 1)
    assert s["pass_percentage"] == 50.0                        # blocked excluded from denominator
    assert s["screenshot_count"] == 2 and s["usage"]["model_calls"] == 0
    assert all(r["publication"]["excel"]["status"] == PUB_PUBLISHED for r in view["results"])
    assert view["has_excel"]

    wb = openpyxl.load_workbook(svc.output_path(run_id))
    ws = wb.worksheets[0]
    col = {c.value: c.column for c in ws[1]}
    assert ws.cell(2, col["Execution Status"]).value == "Passed"
    assert ws.cell(7, col["Execution Status"]).value == "Failed"
    assert "Execution Evidence" in wb.sheetnames

    data, mime, name = svc.export_summary(run_id, "csv")
    assert b"TC-002" in data and name.endswith(".csv")
    payload = json.loads(svc.export_summary(run_id, "json")[0])
    assert any(a["action"] == "run_started" for a in payload["audit"])
    assert "s3cret!" not in json.dumps(payload)


def test_jira_publication_and_duplicate_prevention(svc, jira):
    cases, _ = svc.remote_cases("jira", "PROJ-7")
    run_id = svc.start_run("qa", "jira", cases, cfg(), {"ref": "PROJ-7"}, CREDS)
    view = wait(svc, run_id)
    [r] = view["results"]
    assert r["status"] == "Passed" and r["publication"]["jira"]["status"] == PUB_PUBLISHED
    assert len(jira.comments_by_issue["PROJ-7"]) == 1 and len(jira.attachments) == 1

    svc.retry_publications(run_id, "qa")
    view = wait(svc, run_id)
    assert len(jira.comments_by_issue["PROJ-7"]) == 1           # nothing published twice
    assert view["results"][0]["publication"]["jira"]["status"] == PUB_PUBLISHED
    pubs = svc.store.list_publications(run_id)
    assert len(pubs) == 1


def test_publication_failure_is_separate_from_test_result(svc, jira):
    jira.fail_comment = True
    cases, _ = svc.remote_cases("jira", "PROJ-7")
    run_id = svc.start_run("qa", "jira", cases, cfg(), {"ref": "PROJ-7"}, CREDS)
    view = wait(svc, run_id)
    r = view["results"][0]
    assert r["status"] == "Passed" and r["publication"]["jira"]["status"] == PUB_FAILED
    assert view["summary"]["publication_failed"] == 1 and view["summary"]["passed"] == 1

    jira.fail_comment = False
    svc.retry_publications(run_id, "qa")
    view = wait(svc, run_id)
    assert view["results"][0]["publication"]["jira"]["status"] == PUB_PUBLISHED
    assert view["summary"]["publication_failed"] == 0
    assert len(jira.comments_by_issue["PROJ-7"]) == 1


def test_publishing_turned_off(svc, jira):
    cases, _ = svc.remote_cases("jira", "PROJ-7")
    run_id = svc.start_run("qa", "jira", cases, cfg(publish_results=False), {}, CREDS)
    view = wait(svc, run_id)
    assert view["results"][0]["publication"]["jira"]["status"] == "Skipped"
    assert "PROJ-7" not in jira.comments_by_issue


def test_retry_failed_reruns_only_failed(svc):
    up = svc.save_upload("qa", "cases.xlsx", _xlsx(ROWS))
    sheet = up["sheets"][0]
    cases, meta = svc.excel_cases(up["upload_id"], sheet["name"], sheet["suggested_mapping"])
    run_id = svc.start_run("qa", "excel", cases, cfg(), meta, CREDS)
    wait(svc, run_id)
    before = len(FakeMCPSession.instances)
    assert svc.retry_failed(run_id, "qa", CREDS) == 2          # TC-002 (failed) + TC-003 (blocked)
    view = wait(svc, run_id)
    by = {r["test_key"]: r for r in view["results"]}
    assert by["TC-001"]["retry_count"] == 0 and by["TC-002"]["retry_count"] == 1
    assert len(FakeMCPSession.instances) == before + 1


def test_production_guard_and_missing_base_url(svc, jira, monkeypatch):
    cases, _ = svc.remote_cases("jira", "PROJ-7")
    with pytest.raises(PolicyError):
        svc.start_run("qa", "jira", cases, RunConfig.from_dict({}), {}, CREDS)
    monkeypatch.setenv("TEST_EXECUTION_PRODUCTION_HOSTS", "app.test")
    with pytest.raises(PolicyError, match="production"):
        svc.start_run("qa", "jira", cases, cfg(), {}, CREDS)
    run_id = svc.start_run("qa", "jira", cases, cfg(confirm_production=True), {}, CREDS)
    wait(svc, run_id)


def test_cancel_marks_remaining_not_executed(tmp_path, jira):
    svc = TestExecutionService(str(tmp_path / "d"), str(tmp_path / "db.sqlite"), jira_client_factory=lambda: jira,
                               session_factory=lambda i: FakeMCPSession(slow_tools={"browser_navigate": 0.3}))
    cases = []
    for i in range(5):
        c, _ = svc.remote_cases("jira", "PROJ-7")
        c[0].key = f"PROJ-7#{i}"
        cases += c
    run_id = svc.start_run("qa", "jira", cases, cfg(publish_results=False), {}, CREDS)
    time.sleep(0.4)
    assert svc.cancel(run_id, "qa")
    view = wait(svc, run_id)
    assert view["run"]["status"] == "cancelled"
    assert any(r["status"] == "Not Executed" for r in view["results"])


def test_interrupted_runs_keep_partial_results(tmp_path):
    db = str(tmp_path / "db.sqlite")
    store = ExecutionStore(db)
    store.create_run("TE-20261023-deadbeef", "qa", "excel", "running", {}, {})
    TestExecutionService(str(tmp_path / "d"), db)
    assert store.get_run("TE-20261023-deadbeef")["status"] == RUN_INTERRUPTED


def test_invalid_ids_are_rejected(svc):
    with pytest.raises(PolicyError):
        svc.run_view("../../etc")
    with pytest.raises(PolicyError):
        svc.excel_cases("../x", "S", {"steps": 1})
