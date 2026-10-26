"""
Excel source and result workbook.

Reading supports both layouts the Test Case Generator produces:
  * single cell - one row per test case, steps numbered inside one cell
  * separate rows - a metadata row followed by one row per step with the
    metadata columns left blank (or the ID repeated on every step row)

Writing re-opens the ORIGINAL .xlsx so formatting and unrelated columns and
sheets survive; results go into existing Status/Actual/Evidence columns when
mapped, otherwise into appended result columns.
"""
from __future__ import annotations

import importlib.util
import io
import os
import re
from dataclasses import dataclass, field
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import (
    TestCase, TestStep, TestCaseResult, PASSED, FAILED, BLOCKED, NOT_EXECUTED,
)

# Logical fields -> header spellings seen in real templates (lower-case, compared after normalizing).
FIELD_SYNONYMS: dict[str, list[str]] = {
    "id": ["test case id", "testcase id", "tc id", "tc_id", "test id", "id", "case id", "test case no", "tc no", "key", "test key"],
    "title": ["test case title", "title", "test case name", "name", "summary", "test case", "scenario", "test scenario"],
    "description": ["description", "test description", "objective"],
    "preconditions": ["preconditions", "precondition", "pre-conditions", "pre-requisites", "prerequisites", "pre requisite"],
    "steps": ["test steps", "steps", "step action", "action", "step description", "test step description", "procedure", "steps_list"],
    "step_number": ["test step", "step", "step no", "step number", "step #", "#"],
    "step_expected": ["step expected", "step expected result", "expected step result"],
    "expected": ["expected result", "expected results", "expected", "expected outcome", "expected behaviour", "expected behavior"],
    "test_data": ["test data", "data", "input data", "test input"],
    "actual": ["actual result", "actual results", "actual", "actual outcome"],
    "status": ["status", "execution status", "result", "test result", "pass/fail"],
    "evidence": ["screenshot", "screenshots", "evidence", "screenshot/evidence", "attachment"],
    "defect": ["defect id", "defect", "bug id", "bug"],
    "comments": ["comments", "comment", "remarks", "notes"],
}
REQUIRED_FIELDS = ("steps",)
RESULT_COLUMNS = ["Execution Status", "Actual Result", "Executed At", "Duration (s)", "Failure Details", "Evidence", "Execution Run ID"]
EVIDENCE_SHEET = "Execution Evidence"
SUMMARY_SHEET = "Execution Summary"
MAX_EMBEDDED_IMAGES = int(os.environ.get("TEST_EXECUTION_MAX_EMBEDDED_IMAGES") or "150")
EMBED_WIDTH_PX = 480

_STEP_SPLIT = re.compile(r"(?:^|\n)\s*(?:step\s*)?(\d{1,3})\s*[.):\-]\s+", re.IGNORECASE)
_BULLET_SPLIT = re.compile(r"(?:^|\n)\s*[-*•]\s+")

_STATUS_FILL = {
    PASSED: PatternFill("solid", fgColor="C6EFCE"),
    FAILED: PatternFill("solid", fgColor="FFC7CE"),
    BLOCKED: PatternFill("solid", fgColor="FFEB9C"),
    NOT_EXECUTED: PatternFill("solid", fgColor="E7E6E6"),
}


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").replace("_", " ").strip().lower())


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace("<br>", "\n").strip()


@dataclass
class SheetData:
    name: str
    header_row: int                 # 1-based
    headers: list[str]
    rows: list[tuple[int, list[str]]] = field(default_factory=list)   # (1-based row number, values)


def _load_rows(data: bytes, ext: str) -> list[tuple[str, list[list]]]:
    """[(sheet name, rows as lists)] for .xlsx (openpyxl) or .xls (xlrd)."""
    if ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        return [(s.name, [s.row_values(i) for i in range(s.nrows)]) for s in book.sheets()]
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()


def _detect_header_row(rows: list[list]) -> int:
    """0-based index of the first row that looks like a header (>=2 text cells, one known)."""
    known = {s for syn in FIELD_SYNONYMS.values() for s in syn}
    for i, row in enumerate(rows[:15]):
        cells = [_norm(c) for c in row if _cell_str(c)]
        if len(cells) >= 2 and any(c in known for c in cells):
            return i
    for i, row in enumerate(rows[:15]):
        if sum(1 for c in row if _cell_str(c)) >= 2:
            return i
    return 0


def read_workbook(data: bytes, ext: str) -> list[SheetData]:
    sheets = []
    for name, rows in _load_rows(data, ext):
        if name in (EVIDENCE_SHEET, SUMMARY_SHEET):
            continue
        if not any(any(_cell_str(c) for c in r) for r in rows):
            sheets.append(SheetData(name=name, header_row=1, headers=[]))
            continue
        h = _detect_header_row(rows)
        width = max(len(r) for r in rows)
        headers = [_cell_str(c) for c in rows[h]] + [""] * (width - len(rows[h]))
        body = []
        for idx in range(h + 1, len(rows)):
            vals = [_cell_str(c) for c in rows[idx]] + [""] * (width - len(rows[idx]))
            if any(vals):
                body.append((idx + 1, vals))
        sheets.append(SheetData(name=name, header_row=h + 1, headers=headers, rows=body))
    return sheets


def suggest_mapping(headers: list[str]) -> dict[str, Optional[int]]:
    """Logical field -> column index, best effort. Exact synonym matches win over substring matches."""
    normalized = [_norm(h) for h in headers]
    mapping: dict[str, Optional[int]] = {f: None for f in FIELD_SYNONYMS}
    used: set[int] = set()
    for exact in (True, False):
        for fld, synonyms in FIELD_SYNONYMS.items():
            if mapping[fld] is not None:
                continue
            for syn in synonyms:
                for i, h in enumerate(normalized):
                    if i in used or not h:
                        continue
                    if (h == syn) if exact else (len(syn) > 3 and syn in h):
                        mapping[fld] = i
                        used.add(i)
                        break
                if mapping[fld] is not None:
                    break
    # "Expected Result" next to per-row steps in a separate-rows layout is really the step expectation.
    if mapping["step_number"] is not None and mapping["step_expected"] is None and mapping["expected"] is not None:
        mapping["step_expected"], mapping["expected"] = mapping["expected"], None
    return mapping


def split_numbered(text: str) -> list[tuple[Optional[int], str]]:
    """'1. Open\n2. Click' -> [(1,'Open'),(2,'Click')]. Unnumbered multi-line text -> one item per line."""
    text = (text or "").strip()
    if not text:
        return []
    matches = list(_STEP_SPLIT.finditer(text))
    if len(matches) >= 1 and matches[0].start() <= 1:
        out = []
        for i, m in enumerate(matches):
            end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            body = text[m.end():end].strip()
            if body:
                out.append((int(m.group(1)), " ".join(body.split())))
        if out:
            return out
    if _BULLET_SPLIT.match(text):
        return [(None, " ".join(p.split())) for p in _BULLET_SPLIT.split(text) if p.strip()]
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    return [(None, ln) for ln in lines] if len(lines) > 1 else [(None, " ".join(text.split()))]


def parse_test_cases(sheet: SheetData, mapping: dict[str, Optional[int]]) -> list[TestCase]:
    """Every non-empty row ends up in a test case; problems are attached as issues, never dropped."""
    def col(row: list[str], fld: str) -> str:
        i = mapping.get(fld)
        return row[i] if i is not None and 0 <= i < len(row) else ""

    cases: list[TestCase] = []
    current: Optional[TestCase] = None
    for row_no, row in sheet.rows:
        tc_id, title = col(row, "id"), col(row, "title")
        step_text, step_exp = col(row, "steps"), col(row, "step_expected")
        overall_expected = col(row, "expected")
        if (not step_exp and mapping.get("step_expected") is None and overall_expected
                and sum(1 for n, _ in split_numbered(overall_expected) if n is not None) > 1):
            step_exp, overall_expected = overall_expected, ""   # numbered per-step expectations in one cell
        starts_new = bool(tc_id or title) and (
            current is None
            or (tc_id and tc_id != current.source_ref.get("raw_id"))
            or (not tc_id and title and title != current.title)
        )
        if starts_new:
            current = TestCase(
                key=tc_id or "",
                title=title or tc_id,
                source="excel",
                description=col(row, "description"),
                preconditions=col(row, "preconditions"),
                test_data=col(row, "test_data"),
                expected_result=overall_expected,
                source_ref={"sheet": sheet.name, "row": row_no, "raw_id": tc_id, "step_rows": []},
            )
            cases.append(current)
        elif current is None:
            current = TestCase(key="", title="", source="excel",
                               source_ref={"sheet": sheet.name, "row": row_no, "raw_id": "", "step_rows": []})
            current.issues.append("Rows before the first test case ID/title")
            cases.append(current)
        else:
            # Continuation row: pick up per-step data and late-filled metadata.
            current.test_data = current.test_data or col(row, "test_data")
            current.expected_result = current.expected_result or overall_expected

        if step_text:
            _append_steps(current, step_text, step_exp, col(row, "test_data") if not starts_new else "",
                          row_no, explicit_number=col(row, "step_number"))
            if not starts_new or mapping.get("step_number") is not None:
                current.source_ref["step_rows"].append(row_no)

    _finalize(cases)
    return cases


def _append_steps(tc: TestCase, step_text: str, step_exp: str, data: str, row_no: int, explicit_number: str = ""):
    parts = split_numbered(step_text)
    exp_parts = split_numbered(step_exp) if step_exp else []
    exp_by_num = {n: t for n, t in exp_parts if n is not None}
    for i, (num, text) in enumerate(parts):
        n = len(tc.steps) + 1
        if len(parts) == 1 and explicit_number.isdigit():
            n = int(explicit_number)
        expected = ""
        if num is not None and num in exp_by_num:
            expected = exp_by_num[num]
        elif len(parts) == len(exp_parts):
            expected = exp_parts[i][1]
        elif len(parts) == 1:
            expected = " ".join(step_exp.split())
        if expected.strip().lower() in ("-", "–", "n/a", "na", "none"):
            expected = ""
        tc.steps.append(TestStep(number=n, action_text=text, expected=expected, test_data=data, source_row=row_no))
    if exp_parts and len(parts) > 1 and not exp_by_num and len(parts) != len(exp_parts):
        # Expected text couldn't be aligned with steps: keep it on the last step, not lost.
        tc.steps[-1].expected = (tc.steps[-1].expected + " " + " ".join(step_exp.split())).strip()


def _finalize(cases: list[TestCase]):
    seen: dict[str, int] = {}
    for idx, tc in enumerate(cases, start=1):
        if not tc.key:
            tc.issues.append("Missing Test Case ID")
            tc.key = f"ROW-{tc.source_ref.get('row', idx)}"
        if tc.key in seen:
            seen[tc.key] += 1
            tc.issues.append(f"Duplicate Test Case ID '{tc.key}'")
            tc.key = f"{tc.key}#{seen[tc.key]}"
        else:
            seen[tc.key] = 1
        if not tc.title:
            tc.title = tc.key
        if not tc.steps:
            tc.issues.append("No test steps")
        # A single overall expected result with no per-step expectations belongs on the last step.
        if tc.steps and tc.expected_result and not any(s.expected for s in tc.steps):
            tc.steps[-1].expected = tc.expected_result
        for i, s in enumerate(tc.steps, start=1):
            if not s.number:
                s.number = i


# ── Result workbook ───────────────────────────────────────────────────────────

def _ensure_result_columns(ws, header_row: int, mapping: dict[str, Optional[int]]) -> dict[str, int]:
    """1-based column per result field, reusing mapped Status/Actual/Evidence columns."""
    cols: dict[str, int] = {}
    reuse = {"Execution Status": "status", "Actual Result": "actual", "Evidence": "evidence"}
    existing = {_norm(ws.cell(header_row, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(header_row, c).value}
    next_col = ws.max_column + 1
    header_font = Font(bold=True)
    for name in RESULT_COLUMNS:
        fld = reuse.get(name)
        if fld and mapping.get(fld) is not None:
            cols[name] = mapping[fld] + 1
        elif _norm(name) in existing:
            cols[name] = existing[_norm(name)]
        else:
            cell = ws.cell(header_row, next_col, name)
            cell.font = header_font
            ws.column_dimensions[get_column_letter(next_col)].width = 40 if name in ("Actual Result", "Failure Details") else 20
            cols[name] = next_col
            next_col += 1
    return cols


def _xls_to_workbook(data: bytes) -> openpyxl.Workbook:
    """.xls can't be edited in place; copy the values into a new .xlsx (formatting is not preserved)."""
    import xlrd
    book = xlrd.open_workbook(file_contents=data)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in book.sheets():
        ws = wb.create_sheet(sh.name[:31])
        for r in range(sh.nrows):
            for c, v in enumerate(sh.row_values(r)):
                if v not in ("", None):
                    ws.cell(r + 1, c + 1, v)
    return wb


def build_result_workbook(
    original: bytes, ext: str, sheet_name: str, header_row: int,
    mapping: dict[str, Optional[int]], cases: list[TestCase], results: dict[str, TestCaseResult],
    evidence_reader, evidence_url, embed: bool, summary: dict,
) -> tuple[bytes, list[str]]:
    """
    Returns (xlsx bytes, notes). evidence_reader(evidence) -> bytes|None;
    evidence_url(evidence) -> absolute URL used when not embedding.
    """
    notes: list[str] = []
    wb = _xls_to_workbook(original) if ext == ".xls" else openpyxl.load_workbook(io.BytesIO(original))
    if ext == ".xls":
        notes.append("The source was .xls; results were written to a new .xlsx copy (cell values kept, original formatting not preserved).")
    ws = wb[sheet_name]
    cols = _ensure_result_columns(ws, header_row, mapping)

    image_cls = None
    if embed:
        if importlib.util.find_spec("PIL") is not None:   # openpyxl needs Pillow to embed images
            from openpyxl.drawing.image import Image as _XLImage
            image_cls = _XLImage
        else:
            notes.append("Pillow is not installed, so screenshots are linked instead of embedded.")

    ev_ws = None
    if image_cls or any(r.evidence for r in results.values()):
        if EVIDENCE_SHEET in wb.sheetnames:
            del wb[EVIDENCE_SHEET]
        ev_ws = wb.create_sheet(EVIDENCE_SHEET)
        for c, h in enumerate(["Test Case ID", "Step", "Status", "Captured At", "Screenshot"], start=1):
            ev_ws.cell(1, c, h).font = Font(bold=True)
        ev_ws.column_dimensions["A"].width = 18
        ev_ws.column_dimensions["D"].width = 24
        ev_ws.column_dimensions["E"].width = 70
    ev_row = 2
    embedded = 0
    wrap = Alignment(wrap_text=True, vertical="top")

    for tc in cases:
        res = results.get(tc.key)
        if res is None:
            continue
        row = tc.source_ref.get("row")
        if not row:
            continue
        status_cell = ws.cell(row, cols["Execution Status"], res.status)
        status_cell.fill = _STATUS_FILL.get(res.status, PatternFill())
        ws.cell(row, cols["Actual Result"], res.actual or "").alignment = wrap
        ws.cell(row, cols["Executed At"], res.ended_at or res.started_at or "")
        ws.cell(row, cols["Duration (s)"], round(res.duration_ms / 1000.0, 2) if res.duration_ms else "")
        failure = ""
        if res.failure_reason:
            failure = (f"Step {res.failed_step}: " if res.failed_step else "") + res.failure_reason
        ws.cell(row, cols["Failure Details"], failure).alignment = wrap
        ws.cell(row, cols["Execution Run ID"], res.run_id)

        # Separate-rows layout: per-step status/actual on each step's own row.
        step_rows = set(tc.source_ref.get("step_rows") or [])
        if step_rows:
            for sr, st in zip([s.source_row for s in tc.steps], res.steps):
                if sr in step_rows and sr != row:
                    c = ws.cell(sr, cols["Execution Status"], st.status)
                    c.fill = _STATUS_FILL.get(st.status, PatternFill())
                    ws.cell(sr, cols["Actual Result"], st.actual or st.error or "").alignment = wrap

        if not res.evidence or ev_ws is None:
            continue
        first_ev_row = ev_row
        links = []
        for ev in res.evidence:
            ev_ws.cell(ev_row, 1, tc.key)
            ev_ws.cell(ev_row, 2, "Final" if ev.step_number is None else f"Step {ev.step_number}")
            ev_ws.cell(ev_row, 3, ev.status)
            ev_ws.cell(ev_row, 4, ev.captured_at)
            url = evidence_url(ev)
            placed = False
            if image_cls and embedded < MAX_EMBEDDED_IMAGES:
                raw = evidence_reader(ev)
                if raw:
                    try:
                        img = image_cls(io.BytesIO(raw))
                        if img.width and img.width > EMBED_WIDTH_PX:
                            ratio = EMBED_WIDTH_PX / float(img.width)
                            img.width, img.height = EMBED_WIDTH_PX, int(img.height * ratio)
                        ev_ws.add_image(img, f"E{ev_row}")
                        ev_ws.row_dimensions[ev_row].height = max(20, img.height * 0.75)
                        embedded += 1
                        placed = True
                    except Exception as exc:  # corrupt image: fall back to link
                        notes.append(f"Could not embed screenshot {ev.evidence_id}: {exc}")
            if not placed:
                c = ev_ws.cell(ev_row, 5, url)
                c.hyperlink = url
                c.font = Font(color="0563C1", underline="single")
            links.append(url)
            ev_row += 1
        ev_cell = ws.cell(row, cols["Evidence"], f"{len(res.evidence)} screenshot(s) - see '{EVIDENCE_SHEET}' row {first_ev_row}")
        ev_cell.hyperlink = f"#'{EVIDENCE_SHEET}'!A{first_ev_row}"
        ev_cell.font = Font(color="0563C1", underline="single")

    if embed and image_cls and embedded >= MAX_EMBEDDED_IMAGES:
        notes.append(f"Only the first {MAX_EMBEDDED_IMAGES} screenshots were embedded to keep the workbook small; the rest are linked.")
    if not image_cls and any(r.evidence for r in results.values()):
        notes.append("Screenshot links point at this application and require a logged-in session to open.")

    _write_summary_sheet(wb, summary)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), notes


def _write_summary_sheet(wb, summary: dict):
    if SUMMARY_SHEET in wb.sheetnames:
        del wb[SUMMARY_SHEET]
    ws = wb.create_sheet(SUMMARY_SHEET)
    rows = [
        ("Run ID", summary.get("run_id", "")),
        ("Executed At", summary.get("finished_at") or summary.get("started_at", "")),
        ("Total", summary.get("total", 0)),
        ("Passed", summary.get("passed", 0)),
        ("Failed", summary.get("failed", 0)),
        ("Blocked", summary.get("blocked", 0)),
        ("Not Executed", summary.get("not_executed", 0)),
        ("Pass %", "" if summary.get("pass_percentage") is None else summary["pass_percentage"]),
        ("Pass % formula", summary.get("pass_percentage_formula", "")),
        ("Total Duration (s)", round((summary.get("total_duration_ms") or 0) / 1000.0, 1)),
        ("Screenshots", summary.get("screenshot_count", 0)),
    ]
    for r, (k, v) in enumerate(rows, start=1):
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60
